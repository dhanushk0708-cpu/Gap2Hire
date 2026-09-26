import io
from uuid import uuid4
import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient

from app.core.roles import UserRole
from app.core.security import create_access_token
from app.main import app
from app.models.application import Application
from app.models.candidate import Candidate
from app.models.candidate_source import CandidateSource
from app.models.capability import Capability
from app.models.evidence import Evidence
from app.models.job import Job
from app.models.organization import Organization
from app.models.user import User
from app.models.verification import Verification


def create_in_memory_xlsx(sheets_data: dict[str, list[dict]]) -> bytes:
    """
    Creates an in-memory XLSX workbook with multiple sheets (each sheet is a dataset/concept).
    sheets_data: {"Python": [{"question": "...", "difficulty": "MEDIUM", "type": "CONCEPTUAL"}]}
    """
    wb = openpyxl.Workbook()
    # remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    for sheet_name, rows in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name)
        # Header
        ws.append(["Question", "Concept", "Difficulty", "Question Type", "Expected Topics"])
        for r in rows:
            ws.append([
                r.get("question", ""),
                r.get("concept", sheet_name),
                r.get("difficulty", "MEDIUM"),
                r.get("type", "CONCEPTUAL"),
                r.get("topics", ""),
            ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def setup_test_recruiter_and_job(client: AsyncClient):
    org_id = uuid4()
    user_id = uuid4()
    job_id = uuid4()

    from app.db.session import async_session_factory

    async with async_session_factory() as session:
        org = Organization(id=org_id, name=f"Org {org_id.hex[:6]}", slug=f"org-{org_id.hex[:6]}")
        user = User(
            id=user_id,
            email=f"recruiter_{user_id.hex[:6]}@example.com",
            full_name="Test Recruiter",
            role=UserRole.RECRUITER.value,
            organization_id=org_id,
            password_hash="fakehashpwforunitestsonly",
            is_active=True,
        )
        job = Job(
            id=job_id,
            organization_id=org_id,
            title="Backend Python Developer",
            description="Seeking backend engineer with Python, FastAPI, and PostgreSQL experience.",
            status="PUBLISHED",
            shortlist_size=5,
            created_by=user_id,
        )
        cap1 = Capability(id=uuid4(), job_id=job_id, name="Python", importance="CRITICAL")
        cap2 = Capability(id=uuid4(), job_id=job_id, name="FastAPI", importance="HIGH")
        cap3 = Capability(id=uuid4(), job_id=job_id, name="PostgreSQL", importance="HIGH")
        cap4 = Capability(id=uuid4(), job_id=job_id, name="Docker", importance="MEDIUM")
        cap5 = Capability(id=uuid4(), job_id=job_id, name="Redis", importance="MEDIUM")

        session.add(org)
        await session.flush()
        session.add(user)
        await session.flush()
        session.add(job)
        await session.flush()
        session.add_all([cap1, cap2, cap3, cap4, cap5])
        await session.commit()

    token = create_access_token(user_id=str(user_id))
    headers = {"Authorization": f"Bearer {token}"}
    return headers, org_id, job_id


# ==============================================================================
# PHASE 1: INTERVIEW DATASET UPLOAD & PARSING TESTS
# ==============================================================================

@pytest.mark.asyncio
async def test_valid_3_dataset_xlsx_upload_and_parsing():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        xlsx_bytes = create_in_memory_xlsx({
            "Python": [
                {"question": "Explain Python GIL and memory management.", "difficulty": "HARD", "type": "CONCEPTUAL"},
                {"question": "How do decorators and generators work in Python?", "difficulty": "MEDIUM", "type": "CONCEPTUAL"},
            ],
            "FastAPI": [
                {"question": "How does FastAPI handle dependency injection and async requests?", "difficulty": "HARD", "type": "PRACTICAL"},
                {"question": "Explain Pydantic validation and serialization lifecycle.", "difficulty": "MEDIUM", "type": "CONCEPTUAL"},
            ],
            "PostgreSQL": [
                {"question": "How do you optimize slow PostgreSQL queries with EXPLAIN ANALYZE?", "difficulty": "HARD", "type": "SYSTEM_DESIGN"},
                {"question": "Explain ACID properties and transaction isolation levels in PostgreSQL.", "difficulty": "MEDIUM", "type": "CONCEPTUAL"},
            ],
        })

        resp = await client.post(
            "/api/v1/interview-datasets/upload",
            headers=headers,
            data={"job_id": str(job_id)},
            files={"file": ("interview_questions.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 201
        data = resp.json()
        assert data["filename"] == "interview_questions.xlsx"
        assert data["total_questions"] == 6
        assert len(data["datasets"]) == 3

        concepts = {d["concept"] for d in data["datasets"]}
        assert concepts == {"Python", "FastAPI", "PostgreSQL"}

        # Inspect details endpoint
        file_id = data["file_id"]
        detail_resp = await client.get(f"/api/v1/interview-datasets/{file_id}", headers=headers)
        assert detail_resp.status_code == 200
        detail_data = detail_resp.json()
        assert detail_data["total_questions"] == 6
        assert len(detail_data["datasets"]) == 3
        for d in detail_data["datasets"]:
            assert len(d["questions"]) == 2


@pytest.mark.asyncio
async def test_xlsx_missing_required_column_rejected():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Python"
        # Missing 'Question' column
        ws.append(["Concept", "Difficulty", "Notes"])
        ws.append(["Python", "EASY", "Some note"])
        buf = io.BytesIO()
        wb.save(buf)

        resp = await client.post(
            "/api/v1/interview-datasets/upload",
            headers=headers,
            data={"job_id": str(job_id)},
            files={"file": ("invalid.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 400
        assert "missing required column" in resp.json()["detail"].lower() or "no valid interview questions" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_xlsx_empty_questions_ignored_or_validated():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        xlsx_bytes = create_in_memory_xlsx({
            "Python": [
                {"question": "What are Python metaclasses?", "difficulty": "HARD"},
                {"question": "", "difficulty": "EASY"},  # empty
                {"question": "   ", "difficulty": "MEDIUM"},  # whitespace
                {"question": "How does asyncio event loop work?", "difficulty": "HARD"},
            ]
        })

        resp = await client.post(
            "/api/v1/interview-datasets/upload",
            headers=headers,
            data={"job_id": str(job_id)},
            files={"file": ("partially_empty.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 201
        data = resp.json()
        assert data["total_questions"] == 2  # empty rows skipped


@pytest.mark.asyncio
async def test_xlsx_duplicate_questions_within_dataset_deduplicated():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        xlsx_bytes = create_in_memory_xlsx({
            "FastAPI": [
                {"question": "Explain dependency injection in FastAPI.", "difficulty": "MEDIUM"},
                {"question": "explain dependency injection in fastapi.", "difficulty": "MEDIUM"},  # duplicate
                {"question": "What is the difference between BackgroundTasks and Celery?", "difficulty": "HARD"},
            ]
        })

        resp = await client.post(
            "/api/v1/interview-datasets/upload",
            headers=headers,
            data={"job_id": str(job_id)},
            files={"file": ("duplicates.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 201
        data = resp.json()
        assert data["total_questions"] == 2  # duplicate stripped


@pytest.mark.asyncio
async def test_malformed_corrupted_dataset_rejected():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        resp = await client.post(
            "/api/v1/interview-datasets/upload",
            headers=headers,
            data={"job_id": str(job_id)},
            files={"file": ("corrupt.xlsx", b"not a real excel spreadsheet byte string", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 400
        assert "malformed" in resp.json()["detail"].lower() or "corrupted" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_dataset_tenant_isolation():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers1, org_id1, job_id1 = await setup_test_recruiter_and_job(client)
        headers2, org_id2, job_id2 = await setup_test_recruiter_and_job(client)

        xlsx_bytes = create_in_memory_xlsx({
            "Python": [{"question": "Org 1 specific interview question?", "difficulty": "EASY"}]
        })

        upload_resp = await client.post(
            "/api/v1/interview-datasets/upload",
            headers=headers1,
            data={"job_id": str(job_id1)},
            files={"file": ("org1_questions.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert upload_resp.status_code == 201
        file_id = upload_resp.json()["file_id"]

        # Org 2 attempts to fetch Org 1's dataset -> 404
        cross_resp = await client.get(f"/api/v1/interview-datasets/{file_id}", headers=headers2)
        assert cross_resp.status_code == 404

        # Org 2 list is empty
        list_resp2 = await client.get("/api/v1/interview-datasets", headers=headers2)
        assert list_resp2.status_code == 200
        assert len(list_resp2.json()) == 0


# ==============================================================================
# PHASE 2: CANDIDATE PRE-INTERVIEW ANALYSIS TESTS
# ==============================================================================

async def setup_candidate_with_evidence(
    org_id, job_id, full_name, email,
    demonstrated_caps: list[str],
    claim_caps: list[str],
):
    from app.db.session import async_session_factory
    from sqlalchemy import select

    cand_id = uuid4()
    app_id = uuid4()

    async with async_session_factory() as session:
        cand = Candidate(id=cand_id, full_name=full_name, email=email)
        app_obj = Application(
            id=app_id,
            candidate_id=cand_id,
            job_id=job_id,
            status="SHORTLISTED",
            shortlist_status="SHORTLISTED",
            resume_text=f"Resume of {full_name}. Skills: {', '.join(demonstrated_caps + claim_caps)}.",
        )
        session.add(cand)
        await session.flush()
        session.add(app_obj)
        await session.flush()

        # Query capabilities
        caps = list((await session.scalars(select(Capability).where(Capability.job_id == job_id))).all())
        cap_map = {c.name: c for c in caps}

        # Add demonstrated evidence
        for d_name in demonstrated_caps:
            if d_name in cap_map:
                ev = Evidence(
                    id=uuid4(),
                    application_id=app_id,
                    capability_id=cap_map[d_name].id,
                    source_type="GITHUB",
                    strength="STRONG",
                    provenance="DEMONSTRATED",
                    content=f"Concrete demonstration of {d_name} found in GitHub repo repository.",
                )
                session.add(ev)

        # Add claim evidence
        for c_name in claim_caps:
            if c_name in cap_map:
                ev = Evidence(
                    id=uuid4(),
                    application_id=app_id,
                    capability_id=cap_map[c_name].id,
                    source_type="RESUME",
                    strength="STRONG",
                    provenance="CLAIM",
                    content=f"Candidate stated {c_name} in resume without external proof.",
                )
                session.add(ev)

        await session.commit()
    return app_id


@pytest.mark.asyncio
async def test_pre_analysis_candidate_with_demonstrated_and_claims():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        app_id = await setup_candidate_with_evidence(
            org_id=org_id,
            job_id=job_id,
            full_name="Rahul Sharma",
            email="rahul.sharma@example.com",
            demonstrated_caps=["Python", "FastAPI"],
            claim_caps=["PostgreSQL"],
        )

        resp = await client.get(
            f"/api/v1/applications/{app_id}/interview/rich-pre-analysis",
            headers=headers,
        )
        assert resp.status_code == 200
        data = resp.json()

        assert data["candidate_name"] == "Rahul Sharma"
        assert len(data["candidate_strengths"]) == 2
        strengths_names = {s["capability_name"] for s in data["candidate_strengths"]}
        assert strengths_names == {"Python", "FastAPI"}
        for s in data["candidate_strengths"]:
            assert s["provenance"] == "DEMONSTRATED"

        assert len(data["candidate_claims"]) == 1
        assert data["candidate_claims"][0]["capability_name"] == "PostgreSQL"
        assert data["candidate_claims"][0]["provenance"] == "CLAIM"

        # Docker and Redis should be UNKNOWN
        unknown_names = {u["capability_name"] for u in data["candidate_unknowns"]}
        assert "Docker" in unknown_names
        assert "Redis" in unknown_names

        # Verification targets must include PostgreSQL (CLAIM), Docker (UNKNOWN), Redis (UNKNOWN)
        target_names = {t["capability_name"] for t in data["verification_targets"]}
        assert "PostgreSQL" in target_names
        assert "Docker" in target_names
        assert "Redis" in target_names


@pytest.mark.asyncio
async def test_pre_analysis_preserves_evidence_provenance_and_no_skills_invented():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        app_id = await setup_candidate_with_evidence(
            org_id=org_id,
            job_id=job_id,
            full_name="Claim Only Candidate",
            email="claim.only@example.com",
            demonstrated_caps=[],
            claim_caps=["Python"],
        )

        resp = await client.get(
            f"/api/v1/applications/{app_id}/interview/rich-pre-analysis",
            headers=headers,
        )
        assert resp.status_code == 200
        data = resp.json()

        # No demonstrated strengths
        assert len(data["candidate_strengths"]) == 0
        # Python is strictly CLAIM
        assert len(data["candidate_claims"]) == 1
        assert data["candidate_claims"][0]["capability_name"] == "Python"
        assert data["candidate_claims"][0]["provenance"] == "CLAIM"
        # Summary must state unverified resume claim
        assert "CLAIM" in data["grounded_summary"]


# ==============================================================================
# PHASE 3: AI INTERVIEW PLAN + HR APPROVAL TESTS
# ==============================================================================

@pytest.mark.asyncio
async def test_candidate_specific_interview_plan_generation_and_duplicate_prevention():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        # Upload question dataset
        xlsx_bytes = create_in_memory_xlsx({
            "Python": [
                {"question": "Python Question 1: GIL & Threading?", "difficulty": "HARD"},
                {"question": "Python Question 2: Memory leak profiling?", "difficulty": "MEDIUM"},
                {"question": "Python Question 3: Metaclasses?", "difficulty": "HARD"},
            ],
            "FastAPI": [
                {"question": "FastAPI Question 1: Async dependency injection?", "difficulty": "HARD"},
                {"question": "FastAPI Question 2: Background tasks vs message queues?", "difficulty": "MEDIUM"},
            ],
            "PostgreSQL": [
                {"question": "Postgres Question 1: Index types and write overhead?", "difficulty": "HARD"},
                {"question": "Postgres Question 2: MVCC and connection pooling?", "difficulty": "HARD"},
            ],
        })

        up_resp = await client.post(
            "/api/v1/interview-datasets/upload",
            headers=headers,
            data={"job_id": str(job_id)},
            files={"file": ("questions.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert up_resp.status_code == 201
        file_id = up_resp.json()["file_id"]

        app_id = await setup_candidate_with_evidence(
            org_id=org_id,
            job_id=job_id,
            full_name="Targeted Rahul",
            email="targeted.rahul@example.com",
            demonstrated_caps=["Python"],
            claim_caps=["PostgreSQL", "FastAPI"],
        )

        # Generate AI Interview Plan
        plan_resp = await client.post(
            f"/api/v1/applications/{app_id}/interview/plans",
            headers=headers,
            json={"dataset_file_id": file_id},
        )
        assert plan_resp.status_code == 201
        plan = plan_resp.json()

        assert plan["version"] == 1
        assert plan["status"] == "DRAFT"
        assert len(plan["rounds"]) == 3

        # Check rounds structure
        r1 = plan["rounds"][0]
        assert r1["round_number"] == 1
        assert "Round 1" in r1["title"]
        assert len(r1["questions"]) >= 2

        # Verify question sources are preserved
        all_q_ids = []
        all_q_texts = []
        for r in plan["rounds"]:
            for q in r["questions"]:
                assert q["question_text"]
                assert q["concept"]
                if q["dataset_question_id"]:
                    all_q_ids.append(q["dataset_question_id"])
                all_q_texts.append(q["question_text"].lower())

        # Assert zero duplicate questions across the candidate's plan
        assert len(all_q_texts) == len(set(all_q_texts)), "Duplicate questions found in interview plan!"
        assert len(all_q_ids) == len(set(all_q_ids)), "Duplicate question IDs found in interview plan!"


@pytest.mark.asyncio
async def test_plan_versioning_and_hr_approval():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers, org_id, job_id = await setup_test_recruiter_and_job(client)

        app_id = await setup_candidate_with_evidence(
            org_id=org_id,
            job_id=job_id,
            full_name="Approval Candidate",
            email="approval.candidate@example.com",
            demonstrated_caps=["Python"],
            claim_caps=["FastAPI"],
        )

        # 1. Generate Plan v1
        gen_resp = await client.post(
            f"/api/v1/applications/{app_id}/interview/plans",
            headers=headers,
            json={},
        )
        assert gen_resp.status_code == 201
        v1_plan = gen_resp.json()
        assert v1_plan["version"] == 1
        assert v1_plan["status"] == "DRAFT"
        plan_id = v1_plan["id"]

        # 2. HR Edits Plan v1
        v1_plan["rounds"][0]["estimated_duration_minutes"] = 45
        v1_plan["rounds"][0]["objective"] = "HR customized core fundamentals objective."
        edit_resp = await client.put(
            f"/api/v1/interview-plans/{plan_id}",
            headers=headers,
            json={
                "rounds": v1_plan["rounds"],
                "hr_feedback": "Increased Round 1 duration to 45 mins.",
            },
        )
        assert edit_resp.status_code == 200
        edited = edit_resp.json()
        assert edited["rounds"][0]["estimated_duration_minutes"] == 45
        assert edited["hr_feedback"] == "Increased Round 1 duration to 45 mins."

        # 3. HR Approves Plan v1
        approve_resp = await client.post(
            f"/api/v1/interview-plans/{plan_id}/approve",
            headers=headers,
            json={"notes": "Approved for technical screening round."},
        )
        assert approve_resp.status_code == 200
        approved = approve_resp.json()
        assert approved["status"] == "APPROVED"
        assert approved["approved_at"] is not None

        # 4. Attempt to edit an approved plan must fail (approved plans are immutable)
        bad_edit = await client.put(
            f"/api/v1/interview-plans/{plan_id}",
            headers=headers,
            json={"rounds": approved["rounds"], "hr_feedback": "Illegal edit"},
        )
        assert bad_edit.status_code == 400
        assert "already approved" in bad_edit.json()["detail"].lower()

        # 5. Regenerating plan creates Plan v2 (DRAFT) while preserving historical Plan v1 (APPROVED)
        regen_resp = await client.post(
            f"/api/v1/applications/{app_id}/interview/plans",
            headers=headers,
            json={},
        )
        assert regen_resp.status_code == 201
        v2_plan = regen_resp.json()
        assert v2_plan["version"] == 2
        assert v2_plan["status"] == "DRAFT"

        # List all plans for application -> both v2 and v1 preserved
        list_resp = await client.get(f"/api/v1/applications/{app_id}/interview/plans", headers=headers)
        assert list_resp.status_code == 200
        all_plans = list_resp.json()
        assert len(all_plans) == 2
        versions = [p["version"] for p in all_plans]
        assert 1 in versions and 2 in versions


@pytest.mark.asyncio
async def test_interview_plan_tenant_isolation():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers1, org_id1, job_id1 = await setup_test_recruiter_and_job(client)
        headers2, org_id2, job_id2 = await setup_test_recruiter_and_job(client)

        app_id1 = await setup_candidate_with_evidence(
            org_id=org_id1,
            job_id=job_id1,
            full_name="Tenant Candidate Org1",
            email="tenant.org1@example.com",
            demonstrated_caps=["Python"],
            claim_caps=["PostgreSQL"],
        )

        gen_resp = await client.post(
            f"/api/v1/applications/{app_id1}/interview/plans",
            headers=headers1,
            json={},
        )
        assert gen_resp.status_code == 201
        plan_id = gen_resp.json()["id"]

        # Org 2 attempts to view Org 1's plan -> 404
        cross_get = await client.get(f"/api/v1/interview-plans/{plan_id}", headers=headers2)
        assert cross_get.status_code == 404

        # Org 2 attempts to approve Org 1's plan -> 404
        cross_approve = await client.post(f"/api/v1/interview-plans/{plan_id}/approve", headers=headers2)
        assert cross_approve.status_code == 404
